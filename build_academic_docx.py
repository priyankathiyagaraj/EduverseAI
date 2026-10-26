import os
import sys
from pathlib import Path

# Verify python-docx is available
try:
    import docx
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import parse_xml, OxmlElement
    from docx.oxml.ns import nsdecls, qn
except ImportError:
    print("python-docx not installed in current environment. Please install it.")
    sys.exit(1)

# Check Pillow for diagram generation
try:
    from PIL import Image, ImageDraw, ImageFont
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

BASE_DIR = Path(__file__).resolve().parent
ASSETS_DIR = BASE_DIR / "EDUVERSE_AI_PROJECT_REPORT_ASSETS"
ASSETS_DIR.mkdir(exist_ok=True)


# ==============================================================================
# DIAGRAM GENERATION (PIL)
# ==============================================================================
def create_diagram(filename, width=1200, height=700, bg="#FFFFFF"):
    img = Image.new("RGB", (width, height), bg)
    draw = ImageDraw.Draw(img)
    return img, draw


def draw_rounded_box(draw, xy, text_lines, fill="#EFF6FF", outline="#1E3A8A", text_color="#0F172A", width=2):
    x0, y0, x1, y1 = xy
    draw.rectangle([x0, y0, x1, y1], fill=fill, outline=outline, width=width)
    h_total = len(text_lines) * 22
    start_y = y0 + (y1 - y0 - h_total) // 2
    for i, line in enumerate(text_lines):
        # Center text roughly
        char_w = 8
        txt_w = len(line) * char_w
        txt_x = x0 + (x1 - x0 - txt_w) // 2
        draw.text((max(x0 + 5, txt_x), start_y + i * 22), line, fill=text_color)


def draw_arrow(draw, start, end, color="#1E3A8A", width=2):
    x0, y0 = start
    x1, y1 = end
    draw.line([x0, y0, x1, y1], fill=color, width=width)
    # arrow head
    if y0 == y1: # horizontal
        direction = 1 if x1 > x0 else -1
        draw.polygon([(x1, y1), (x1 - 10 * direction, y1 - 6), (x1 - 10 * direction, y1 + 6)], fill=color)
    elif x0 == x1: # vertical
        direction = 1 if y1 > y0 else -1
        draw.polygon([(x1, y1), (x1 - 6, y1 - 10 * direction), (x1 + 6, y1 - 10 * direction)], fill=color)


def generate_all_diagrams():
    if not HAS_PIL:
        print("Pillow not installed; skipping graphic diagram rendering.")
        return

    # 1. Architecture Diagram
    img, draw = create_diagram("fig_4_1_architecture.png", 1200, 750)
    draw.text((380, 20), "EDUVERSE AI - MULTI-TIER SYSTEM ARCHITECTURE", fill="#1E3A8A")
    # Presentation Layer
    draw_rounded_box(draw, (50, 70, 1150, 170), ["PRESENTATION TIER (CLIENT BROWSER)", "Jinja2 HTML5 Views | Custom SaaS CSS3 (#1E3A8A / #2563EB) | Vanilla JS & Chart.js"], fill="#F0F9FF", outline="#0284C7")
    draw_arrow(draw, (600, 170), (600, 230))
    # Flask Application Layer
    draw_rounded_box(draw, (50, 230, 1150, 360), [
        "APPLICATION ROUTING & CONTROLLER TIER (FLASK 3.1.0)",
        "Application Factory (create_app) | CSRF Protection | Session Authentication",
        "Modular Blueprints: main_bp (/), auth_bp (/auth), colleges_bp (/colleges), comparison_bp (/compare)",
        "recommendations_bp (/recommend), analytics_bp (/analytics), admin_bp (/admin), student_bp (/student)"
    ], fill="#EFF6FF", outline="#2563EB")
    draw_arrow(draw, (600, 360), (600, 420))
    # Business Logic Layer
    draw_rounded_box(draw, (50, 420, 1150, 540), [
        "BUSINESS LOGIC & SERVICE LAYER",
        "CollegeService (Faceted Query Search) | AssessmentService (12th Marks & TNEA Cutoff Math)",
        "RecommendationService (Content-Based Multi-Criteria Scoring Engine: 35% Course, 25% Loc, 20% Budget, 10% Type, 10% Fac)",
        "AICounselorService (Google Gemini 1.5 Flash API + Autonomous Rule-Based Fallback Engine)"
    ], fill="#F8FAFC", outline="#475569")
    draw_arrow(draw, (600, 540), (600, 600))
    # Database Layer
    draw_rounded_box(draw, (50, 600, 1150, 700), [
        "DATA PERSISTENCE & ORM TIER (SQLALCHEMY 2.0)",
        "Relational Store: SQLite (instance/eduverse_dev.db) / MySQL via PyMySQL Driver",
        "1,129 Colleges (All 38 Districts) | 6,290 Course Offerings | Verified Official Websites | Geocoordinates"
    ], fill="#ECFDF5", outline="#059669")
    img.save(ASSETS_DIR / "fig_4_1_architecture.png")

    # 2. Use Case Diagram
    img, draw = create_diagram("fig_4_3_use_case.png", 1100, 700)
    draw.text((360, 20), "UML USE CASE DIAGRAM - EDUVERSE AI", fill="#1E3A8A")
    # System Boundary
    draw.rectangle([250, 60, 850, 660], outline="#2563EB", width=2, fill="#F8FAFC")
    draw.text((450, 75), "<< System >> EduVerse AI", fill="#1E3A8A")
    # Student Actor
    draw_rounded_box(draw, (30, 250, 170, 330), ["STUDENT", "ACTOR"], fill="#EFF6FF", outline="#1E3A8A")
    # Admin Actor
    draw_rounded_box(draw, (30, 480, 170, 560), ["ADMIN", "ACTOR"], fill="#EFF6FF", outline="#1E3A8A")
    # External Gemini Actor
    draw_rounded_box(draw, (930, 250, 1070, 330), ["GOOGLE", "GEMINI AI"], fill="#FEF3C7", outline="#D97706")
    # Use cases
    draw_rounded_box(draw, (320, 110, 560, 160), ["Browse & Search Colleges"], fill="#FFFFFF", outline="#0284C7")
    draw_rounded_box(draw, (320, 180, 560, 230), ["Filter by District & Budget"], fill="#FFFFFF", outline="#0284C7")
    draw_rounded_box(draw, (320, 250, 560, 300), ["Compare Colleges Side-by-Side"], fill="#FFFFFF", outline="#0284C7")
    draw_rounded_box(draw, (320, 320, 560, 370), ["Compute 12th TNEA Cutoff"], fill="#FFFFFF", outline="#0284C7")
    draw_rounded_box(draw, (320, 390, 560, 440), ["Receive Recommendations"], fill="#FFFFFF", outline="#0284C7")
    draw_rounded_box(draw, (320, 460, 560, 510), ["Chat with AI Counselor"], fill="#FFFFFF", outline="#0284C7")
    draw_rounded_box(draw, (600, 460, 800, 510), ["<<Fallback>> Rule Engine"], fill="#F0FDF4", outline="#16A34A")
    draw_rounded_box(draw, (320, 530, 560, 580), ["Manage Colleges (CRUD)"], fill="#FFFFFF", outline="#DC2626")
    draw_rounded_box(draw, (320, 600, 560, 650), ["Triage Student Queries"], fill="#FFFFFF", outline="#DC2626")
    # Connections
    draw.line([(170, 290), (320, 135)], fill="#1E3A8A", width=2)
    draw.line([(170, 290), (320, 205)], fill="#1E3A8A", width=2)
    draw.line([(170, 290), (320, 275)], fill="#1E3A8A", width=2)
    draw.line([(170, 290), (320, 345)], fill="#1E3A8A", width=2)
    draw.line([(170, 290), (320, 415)], fill="#1E3A8A", width=2)
    draw.line([(170, 290), (320, 485)], fill="#1E3A8A", width=2)
    draw.line([(170, 520), (320, 555)], fill="#1E3A8A", width=2)
    draw.line([(170, 520), (320, 625)], fill="#1E3A8A", width=2)
    draw.line([(560, 485), (930, 290)], fill="#D97706", width=2)
    draw.line([(560, 485), (600, 485)], fill="#16A34A", width=2)
    img.save(ASSETS_DIR / "fig_4_3_use_case.png")

    # 3. DFD Level 0 (Context Diagram)
    img, draw = create_diagram("fig_4_4_dfd_context.png", 1000, 550)
    draw.text((320, 20), "DATA FLOW DIAGRAM (LEVEL 0 - CONTEXT DIAGRAM)", fill="#1E3A8A")
    draw_rounded_box(draw, (50, 200, 220, 320), ["STUDENT USER", "(External Entity)"], fill="#EFF6FF", outline="#1E3A8A")
    draw_rounded_box(draw, (380, 170, 620, 350), ["0.0", "EDUVERSE AI PLATFORM", "(Central System Process)"], fill="#DBEAFE", outline="#1D4ED8")
    draw_rounded_box(draw, (780, 200, 950, 320), ["ADMINISTRATOR", "(External Entity)"], fill="#EFF6FF", outline="#1E3A8A")
    draw_rounded_box(draw, (400, 440, 600, 510), ["GOOGLE GEMINI API", "(External Cloud Service)"], fill="#FEF3C7", outline="#D97706")
    # Flows
    draw_arrow(draw, (220, 230), (380, 230))
    draw.text((230, 210), "Marks, Search, Prefs", fill="#0F172A")
    draw_arrow(draw, (380, 280), (220, 280))
    draw.text((230, 290), "Colleges, Cutoffs, Links", fill="#0F172A")
    draw_arrow(draw, (780, 250), (620, 250))
    draw.text((640, 230), "College CRUD, Edits", fill="#0F172A")
    draw_arrow(draw, (500, 350), (500, 440))
    draw.text((510, 380), "Student Prompt", fill="#0F172A")
    draw_arrow(draw, (480, 440), (480, 350))
    draw.text((370, 380), "AI Guidance", fill="#0F172A")
    img.save(ASSETS_DIR / "fig_4_4_dfd_context.png")

    # 4. Entity-Relationship Diagram
    img, draw = create_diagram("fig_4_7_er_diagram.png", 1100, 650)
    draw.text((360, 20), "ENTITY-RELATIONSHIP (ER) DIAGRAM - DATABASE TOPOLOGY", fill="#1E3A8A")
    draw_rounded_box(draw, (50, 80, 250, 220), ["districts", "PK: id", "name (UNIQUE)", "code, zone"], fill="#FEF2F2", outline="#DC2626")
    draw_rounded_box(draw, (380, 80, 720, 320), [
        "colleges (1,129 Records)",
        "PK: id",
        "FK: district_id -> districts.id",
        "name, code, district, city",
        "college_type, autonomous_status",
        "official_website, website_status",
        "latitude, longitude, rating, accreditation"
    ], fill="#EFF6FF", outline="#1D4ED8")
    draw_rounded_box(draw, (850, 80, 1050, 240), ["courses (6,290 Records)", "PK: id", "FK: college_id", "name, duration_years", "annual_fee, cutoff_marks"], fill="#F0FDF4", outline="#16A34A")
    draw_rounded_box(draw, (850, 280, 1050, 400), ["facilities", "PK: id", "FK: college_id", "facility_name"], fill="#F0FDF4", outline="#16A34A")
    draw_rounded_box(draw, (380, 420, 620, 580), ["users", "PK: id", "name, email (UNIQUE)", "password_hash", "role ('student'/'admin')"], fill="#FAF5FF", outline="#7C3AED")
    draw_rounded_box(draw, (750, 460, 1050, 600), ["student_profiles", "PK: id", "FK: user_id -> users.id", "phone, school, stream", "assessment_result (JSON)"], fill="#FFFBEB", outline="#D97706")
    # Connections
    draw_arrow(draw, (250, 150), (380, 150))
    draw.text((270, 130), "1 : N", fill="#B91C1C")
    draw_arrow(draw, (720, 150), (850, 150))
    draw.text((740, 130), "1 : N", fill="#15803D")
    draw_arrow(draw, (720, 250), (850, 330))
    draw.text((740, 270), "1 : N", fill="#15803D")
    draw_arrow(draw, (620, 500), (750, 500))
    draw.text((640, 480), "1 : 1", fill="#B45309")
    img.save(ASSETS_DIR / "fig_4_7_er_diagram.png")
    print("All architecture and system diagrams rendered successfully in ASSETS_DIR.")


# ==============================================================================
# DOCX BUILDER WITH ACADEMIC FORMATTING
# ==============================================================================
def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)


def build_academic_docx():
    print("Initializing Microsoft Word Document...")
    doc = Document()

    # 1. Page Setup: A4, Left 1.5 in, Right 1.0 in, Top 1.0 in, Bottom 1.0 in
    for sec in doc.sections:
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        sec.left_margin = Inches(1.5)
        sec.right_margin = Inches(1.0)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        # Header & Footer setup
        header = sec.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("EDUVERSE AI – PROJECT REPORT")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)

        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Department of Computer Applications | Anna University")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(100, 116, 139)

    # 2. Configure Standard Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)

    # Heading 1
    h1 = doc.styles['Heading 1']
    h1.font.name = 'Times New Roman'
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(30, 58, 138) # Dark Blue #1E3A8A
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Heading 2
    h2 = doc.styles['Heading 2']
    h2.font.name = 'Times New Roman'
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(37, 99, 235) # Blue #2563EB
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    # Heading 3
    h3 = doc.styles['Heading 3']
    h3.font.name = 'Times New Roman'
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(15, 23, 42)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)

    # ==========================================================================
    # PRELIMINARY PAGES
    # ==========================================================================
    # 1. Cover Page
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c_title = cp.add_run("EDUVERSE AI\nDISCOVER, COMPARE & CHOOSE YOUR COLLEGE IN TAMIL NADU\n\n")
    c_title.font.size = Pt(18)
    c_title.font.bold = True
    c_title.font.color.rgb = RGBColor(30, 58, 138)

    c_sub = cp.add_run("PROJECT REPORT\n\nSubmitted by\n\n")
    c_sub.font.size = Pt(13)
    c_sub.font.bold = True

    c_stu = cp.add_run("R. PRIYADHARSHINI (Reg. No: 9525621125)\nT. PRIYANKA (Reg. No: 9525621126)\n\n")
    c_stu.font.size = Pt(13)
    c_stu.font.bold = True
    c_stu.font.color.rgb = RGBColor(37, 99, 235)

    c_deg = cp.add_run("In partial fulfilment for the award of the degree of\nMASTER OF COMPUTER APPLICATIONS\n\n\n[COLLEGE EMBLEM / LOGO PLACEHOLDER]\n\n\n")
    c_deg.font.size = Pt(12)

    c_inst = cp.add_run("DEPARTMENT OF MASTER OF COMPUTER APPLICATIONS\n[COLLEGE NAME PLACEHOLDER]\n(AN AUTONOMOUS INSTITUTION) – AFFILIATED TO ANNA UNIVERSITY, CHENNAI\n[LOCATION, TAMIL NADU - PIN CODE]\n\nOCTOBER 2026\n")
    c_inst.font.size = Pt(12)
    c_inst.font.bold = True

    doc.add_page_break()

    # 2. Bonafide Certificate
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BONAFIDE CERTIFICATE\n")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    p_cert = doc.add_paragraph(
        "Certified that this Project Report titled \"EDUVERSE AI – DISCOVER, COMPARE & CHOOSE YOUR COLLEGE IN TAMIL NADU\" "
        "is the bonafide work of R. PRIYADHARSHINI (Roll No: 9525621125) and T. PRIYANKA (Roll No: 9525621126) "
        "who carried out the work under my supervision. Certified further that to the best of my knowledge the work reported herein "
        "does not form part of any other thesis or dissertation on the basis of which a degree or award was conferred on an earlier occasion on this or any other candidate.\n\n\n"
    )
    p_cert.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Signature Table
    sig_tab = doc.add_table(rows=1, cols=2)
    sig_tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1, c2 = sig_tab.rows[0].cells
    c1.paragraphs[0].add_run("[INTERNAL GUIDE NAME]\nInternal Guide & Assistant Professor\nDepartment of Computer Applications\n[College Name]\n[Location, Tamil Nadu]")
    c2.paragraphs[0].add_run("[HEAD OF THE DEPARTMENT]\nProfessor and Head of the Department\nDepartment of Computer Applications\n[College Name]\n[Location, Tamil Nadu]")
    c1.paragraphs[0].runs[0].font.name = "Times New Roman"
    c2.paragraphs[0].runs[0].font.name = "Times New Roman"

    doc.add_page_break()

    # 3. Certificate of Viva-Voce
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CERTIFICATE OF VIVA-VOCE EXAMINATION\n")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    p_viva = doc.add_paragraph(
        "This is to certify that the Viva-Voce Examination of the Project Work entitled "
        "\"EDUVERSE AI – DISCOVER, COMPARE & CHOOSE YOUR COLLEGE IN TAMIL NADU\" submitted by "
        "Ms. R. PRIYADHARSHINI (Reg. No. 9525621125) and Ms. T. PRIYANKA (Reg. No. 9525621126) in partial fulfillment of "
        "the requirements for the award of the degree of Master of Computer Applications (MCA) was conducted on ______________ "
        "in the Department of Master of Computer Applications, [College Name], [Location, Tamil Nadu].\n\n\n"
    )
    p_viva.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    viva_tab = doc.add_table(rows=1, cols=2)
    viva_tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    vc1, vc2 = viva_tab.rows[0].cells
    vc1.paragraphs[0].add_run("INTERNAL EXAMINER\nName:\nDesignation:\nAddress:\nDate:")
    vc2.paragraphs[0].add_run("EXTERNAL EXAMINER\nName:\nDesignation:\nAddress:\nDate:")

    doc.add_page_break()

    # 4. Declaration
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DECLARATION\n")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    p_dec = doc.add_paragraph(
        "We, R. PRIYADHARSHINI (Register No: 9525621125) and T. PRIYANKA (Register No: 9525621126), hereby declare that "
        "the project report entitled \"EDUVERSE AI – DISCOVER, COMPARE & CHOOSE YOUR COLLEGE IN TAMIL NADU\" submitted to "
        "the Department of Computer Applications, [College Name], affiliated to Anna University, Chennai, is a record of original "
        "work done by us under the supervision of [INTERNAL GUIDE NAME].\n\n"
        "We further declare that to the best of our knowledge, the work reported herein does not form part of any other thesis or "
        "project on the basis of which a degree has been conferred previously.\n\n"
    )
    p_dec.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    dec_tab = doc.add_table(rows=1, cols=2)
    dec_tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    dc1, dc2 = dec_tab.rows[0].cells
    dc1.paragraphs[0].add_run("Place: [Location]\nDate:")
    dc2.paragraphs[0].add_run("R. PRIYADHARSHINI (9525621125)\nT. PRIYANKA (9525621126)\nCandidates")

    doc.add_page_break()

    # 5. Abstract (English)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ABSTRACT\n")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    p_abs = doc.add_paragraph(
        "Choosing an appropriate higher education institution and undergraduate program represents one of the most critical and challenging life decisions "
        "faced by students upon completing their higher secondary education (12th standard) in Tamil Nadu. The higher education ecosystem across the state "
        "encompasses more than a thousand colleges distributed across 38 administrative districts, featuring diverse institutional types, differing university "
        "affiliations, variable tuition fee structures, and specialized accreditation statuses. In the absence of a unified, objective, and student-friendly discovery "
        "mechanism, students and parents frequently encounter fragmented, unverified, or commercially biased information across third-party portals, leading to "
        "suboptimal course selection, confusion regarding Tamil Nadu Engineering Admissions (TNEA) cutoff calculations, and missed higher education opportunities.\n\n"
        "To address these systemic challenges, this project presents EduVerse AI, a full-stack, AI-augmented academic discovery, comparison, and decision-support web "
        "platform engineered explicitly for Tamil Nadu higher education. The system is developed using Python Flask as the application framework, SQLAlchemy "
        "Object-Relational Mapper (ORM), a structured database, and a responsive frontend built with modern HTML5, CSS3, JavaScript, and Jinja2 templating. "
        "EduVerse AI integrates an empirical catalog of 1,129 verified colleges and 6,290 courses spanning all 38 districts of Tamil Nadu. Each indexed institution is "
        "mapped with verified institutional parameters, geolocation coordinates, and its authenticated official website URL.\n\n"
        "The platform provides an end-to-end user experience comprising multi-attribute search and filtering (by district, college type, autonomous status, course discipline, "
        "and budget ceiling), a side-by-side college comparison matrix, dynamic district-level analytics powered by Chart.js, and an academic assessment engine that "
        "computes 12th standard aggregate percentages and TNEA engineering cutoffs. A content-based recommendation service evaluates student preferences using a weighted "
        "multi-criteria scoring algorithm to generate ranked institutional recommendations with explainable factor diagnostics. Furthermore, the platform integrates an AI "
        "Education Counselor powered by Google Gemini API with an automated rule-based fallback guidance engine, enabling students to receive conversational guidance on "
        "course suitability, career opportunities, and admission procedures. Empirical validation of the system confirms that all 24 live application endpoints respond with "
        "HTTP 200 OK, and an automated regression test suite comprising 85 pytest test cases passed with a 100% success rate. The outcome is a scalable, transparent, "
        "and student-centered educational decision-support ecosystem."
    )
    p_abs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_page_break()

    # 6. Tamil Abstract
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("சுருக்கம் (TAMIL ABSTRACT)\n")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    p_tabs = doc.add_paragraph(
        "பன்னிரண்டாம் வகுப்பு பொதுத்தேர்வை முடித்த மாணவர்கள் தங்களின் அடுத்தகட்ட உயர்கல்வி மற்றும் கல்லூரிப் பிரிவை தேர்ந்தெடுப்பது என்பது அவர்களின் எதிர்கால "
        "வாழ்க்கையை தீர்மானிக்கும் மிக முக்கியமான முடிவாகும். தமிழ்நாட்டில் உள்ள 38 மாவட்டங்களில் ஆயிரத்திற்கும் மேற்பட்ட அரசு, அரசு உதவிபெறும், தனியார் சுயநிதி "
        "மற்றும் தன்னாட்சிப் (Autonomous) பொறியியல், கலை மற்றும் அறிவியல் கல்லூரிகள் இயங்கி வருகின்றன. இத்தகைய பரந்த கல்விச் சூழலில், மாணவர்கள் தங்களின் மதிப்பெண்கள், "
        "ஆர்வங்கள் மற்றும் இருப்பிடத்திற்கு ஏற்ற சிறந்த கல்லூரிகளை கண்டறிவதில் போதிய விழிப்புணர்வின்மையும், சரியான வழிகாட்டலின்மையும் பெரும் சவாலாக உள்ளது. பல வணிக "
        "நோக்குடைய இணையதளங்கள் முழுமையற்ற மற்றும் விளம்பரச் சார்புடைய தகவல்களையே வழங்குகின்றன.\n\n"
        "இச்சிக்கல்களுக்குத் தீர்வாக, EduVerse AI என்ற முழுமையான, செயற்கை நுண்ணறிவு (Artificial Intelligence) சார்ந்த உயர்கல்வி வழிகாட்டுதல் மற்றும் கல்லூரி ஒப்பீட்டு "
        "வலைப்பின்னல் தளம் உருவாக்கப்பட்டுள்ளது. இத்தளம் பைதான் பிளாஸ்க் (Python Flask) கட்டமைப்பிலும், SQL தரவுத்தள (Database) அடிப்படையிலும், HTML, CSS மற்றும் "
        "ஜாவாஸ்கிரிப்ட் (JavaScript) தொழில் நுட்பங்களுடனும் கட்டமைக்கப்பட்டுள்ளது. இதில் தமிழ்நாட்டின் அனைத்து 38 மாவட்டங்களிலும் உள்ள 1,129 அங்கீகரிக்கப்பட்ட கல்லூரிகள் "
        "மற்றும் 6,290 பாடப்பிரிவுகளின் அதிகாரப்பூர்வ விபரங்கள், அமைவிடம் மற்றும் சரிபார்க்கப்பட்ட அதிகாரப்பூர்வ இணையதள முகவரிகள் (Official Websites) ஒருங்கிணைக்கப்பட்டுள்ளன.\n\n"
        "மாணவர்கள் தங்கள் மாவட்டங்கள், கட்டண வரம்புகள், தன்னாட்சி அந்தஸ்து மற்றும் பாடப்பிரிவுகளின் அடிப்படையில் கல்லூரிகளை தேடவும் (Search & Filter), பல கல்லூரிகளை ஒரே "
        "திரையில் ஒப்பிட்டுப் பார்க்கவும் (College Comparison), மாவட்ட வாரியான கல்வி புள்ளிவிபரங்களை (Analytics) காணவும் முடியும். மேலும், 12-ஆம் வகுப்பு மதிப்பெண்களை "
        "உள்ளீடு செய்து TNEA பொறியியல் கலந்தாய்வு Cutoff மதிப்பெண்களை கணக்கிடும் முறையும், விருப்பங்களின் அடிப்படையில் கல்லூரிகளை வரிசைப்படுத்தி பரிந்துரைக்கும் "
        "(Recommendation Algorithm) அமைப்பும் இதில் நிறுவப்பட்டுள்ளது. கூகிள் ஜெமினி (Google Gemini API) மற்றும் விதிமுறை சார்ந்த ஆலோசனை இயந்திரம் (Rule-Based Counselor "
        "Engine) மூலம் மாணவர்கள் தங்களுக்குப் பொருத்தமான உயர்கல்வி மற்றும் வேலைவாய்ப்பு ஆலோசனைகளை நேரடியாகப் பெற முடிகிறது. இத்திட்டத்தின் செயல்முறைகள் 85 தானியங்கி "
        "சோதனைகள் (Pytest) மூலம் வெற்றிகரமாக சரிபார்க்கப்பட்டு 100% துல்லியத்துடன் இயங்குவது உறுதி செய்யப்பட்டுள்ளது."
    )
    p_tabs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_page_break()

    # 7. Acknowledgement
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ACKNOWLEDGEMENT\n")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    p_ack = doc.add_paragraph(
        "We thank the Almighty God for His divine grace, strength, and blessings throughout the successful completion of this project.\n\n"
        "We express our sincere gratitude to the respected Management and our Principal, [Principal Name], [College Name], for providing state-of-the-art "
        "computing laboratories and academic infrastructure to carry out this project work.\n\n"
        "We convey our deep sense of gratitude to [Head of Department Name], Professor and Head of the Department of Master of Computer Applications, "
        "for the constant encouragement and leadership extended to us throughout our postgraduate studies.\n\n"
        "We are profoundly indebted to our Internal Guide, [Internal Guide Name], Assistant Professor, Department of Master of Computer Applications, "
        "for the invaluable guidance, constructive suggestions, and technical critique provided during the course of this project.\n\n"
        "We extend our sincere thanks to our Project Coordinator, [Project Coordinator Name], and all faculty members of the Department of Computer Applications "
        "for their timely support and co-operation.\n\n"
        "Finally, we express our heartfelt love and thanks to our parents, family members, and friends whose unconditional sacrifices and encouragement "
        "have been our greatest strength.\n\n\n"
        "R. PRIYADHARSHINI (9525621125)\n"
        "T. PRIYANKA (9525621126)"
    )
    p_ack.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_page_break()

    # 8. Table of Contents & Lists
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TABLE OF CONTENTS\n")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    toc_items = [
        ("BONAFIDE CERTIFICATE", "ii"),
        ("CERTIFICATE OF VIVA-VOCE EXAMINATION", "iii"),
        ("DECLARATION", "iv"),
        ("ABSTRACT", "v"),
        ("சுருக்கம் (TAMIL ABSTRACT)", "vi"),
        ("ACKNOWLEDGEMENT", "vii"),
        ("LIST OF TABLES", "x"),
        ("LIST OF FIGURES", "xi"),
        ("LIST OF ABBREVIATIONS", "xii"),
        ("CHAPTER 1: INTRODUCTION", "1"),
        ("  1.1 Background of the Study", "1"),
        ("  1.2 Problem Statement", "4"),
        ("  1.3 Objectives of the Project", "6"),
        ("  1.4 Scope of the Project", "7"),
        ("  1.5 Organisation of the Report", "8"),
        ("CHAPTER 2: LITERATURE REVIEW", "9"),
        ("  2.1 Evolution of Online College Discovery Systems", "9"),
        ("  2.2 Education Recommendation Systems", "10"),
        ("  2.3 Artificial Intelligence in Student Decision Support", "11"),
        ("  2.4 College Search and Comparison Platforms", "13"),
        ("  2.5 Web-Based Information Retrieval", "14"),
        ("  2.6 Research Gaps and Project Motivation", "15"),
        ("CHAPTER 3: SYSTEM ANALYSIS", "17"),
        ("  3.1 Existing System", "17"),
        ("  3.2 Proposed System", "18"),
        ("  3.3 Functional Requirements", "20"),
        ("  3.4 Non-Functional Requirements", "22"),
        ("  3.5 Hardware Requirements", "23"),
        ("  3.6 Software Requirements", "24"),
        ("  3.7 Feasibility Study", "25"),
        ("  3.8 System Constraints", "27"),
        ("CHAPTER 4: SYSTEM DESIGN", "29"),
        ("  4.1 Overall System Architecture", "29"),
        ("  4.2 Layered Architectural Design", "31"),
        ("  4.3 Use Case Diagram", "34"),
        ("  4.4 Data Flow Diagrams (Context & Level-1)", "36"),
        ("  4.5 Application Workflow", "38"),
        ("  4.6 Database Design (ER Model & Schema)", "40"),
        ("  4.7 API Design", "45"),
        ("  4.8 User Interface Design", "48"),
        ("  4.9 Design Principles", "50"),
        ("CHAPTER 5: SYSTEM IMPLEMENTATION", "52"),
        ("  5.1 Development Environment", "52"),
        ("  5.2 Project Directory Structure", "53"),
        ("  5.3 Flask Application Initialisation", "56"),
        ("  5.4 Frontend Implementation", "58"),
        ("  5.5 College Data Implementation", "60"),
        ("  5.6 College Search and Filtering Implementation", "62"),
        ("  5.7 College Comparison Implementation", "64"),
        ("  5.8 AI and Recommendation Implementation", "66"),
        ("  5.9 Analytics Implementation", "70"),
        ("  5.10 Error Handling Strategy", "72"),
        ("  5.11 Security Implementation", "73"),
        ("  5.12 Important Code Snippets", "75"),
        ("CHAPTER 6: TESTING AND PERFORMANCE EVALUATION", "80"),
        ("  6.1 Testing Strategy", "80"),
        ("  6.2 Automated Unit and Integration Test Cases", "82"),
        ("  6.3 API Endpoint Testing", "85"),
        ("  6.4 Frontend and User Interface Testing", "87"),
        ("  6.5 Performance Evaluation", "89"),
        ("  6.6 Security Audit and Testing", "90"),
        ("  6.7 Results and Discussion", "92"),
        ("  6.8 Application Screenshots and Interface Analysis", "94"),
        ("CHAPTER 7: CONCLUSION AND FUTURE ENHANCEMENTS", "96"),
        ("  7.1 Conclusion", "96"),
        ("  7.2 System Limitations", "97"),
        ("  7.3 Future Enhancements", "98"),
        ("APPENDICES (A to H)", "101"),
        ("REFERENCES", "123")
    ]

    t_toc = doc.add_table(rows=len(toc_items), cols=2)
    t_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (title, page) in enumerate(toc_items):
        r_title = t_toc.rows[idx].cells[0].paragraphs[0].add_run(title)
        r_page = t_toc.rows[idx].cells[1].paragraphs[0].add_run(page)
        t_toc.rows[idx].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_title.font.name = "Times New Roman"
        r_page.font.name = "Times New Roman"
        if "CHAPTER" in title or "APPENDIX" in title or "REFERENCES" in title:
            r_title.font.bold = True
            r_page.font.bold = True

    doc.add_page_break()

    # List of Tables
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LIST OF TABLES\n")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    tables_list = [
        ("3.1", "Minimum and Recommended Hardware Specifications", "23"),
        ("3.2", "Software Environment and Framework Dependencies", "24"),
        ("4.1", "Database Entity Specification: colleges Table", "41"),
        ("4.2", "Database Entity Specification: courses Table", "42"),
        ("4.3", "Database Entity Specification: districts Table", "43"),
        ("4.4", "Database Entity Specification: facilities Table", "43"),
        ("4.5", "Database Entity Specification: users Table", "44"),
        ("4.6", "Comprehensive API Endpoint Catalog", "46"),
        ("5.1", "Multi-Criteria Recommendation Scoring Weights", "68"),
        ("6.1", "Test Suite Execution Summary (pytest)", "83"),
        ("6.2", "Live API Endpoint Verification and Response Times", "86"),
        ("6.3", "Browser Cross-Platform Verification Matrix", "88"),
        ("6.4", "System Latency and Performance Benchmark", "89"),
        ("6.5", "Security Vulnerability Assessment Matrix", "91")
    ]
    t_lt = doc.add_table(rows=len(tables_list) + 1, cols=3)
    t_lt.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_row = t_lt.rows[0].cells
    h_row[0].paragraphs[0].add_run("Table No.").font.bold = True
    h_row[1].paragraphs[0].add_run("Title").font.bold = True
    h_row[2].paragraphs[0].add_run("Page No.").font.bold = True
    for i, (t_no, t_title, t_pg) in enumerate(tables_list):
        row = t_lt.rows[i + 1].cells
        row[0].paragraphs[0].add_run(t_no)
        row[1].paragraphs[0].add_run(t_title)
        row[2].paragraphs[0].add_run(t_pg)
        row[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_page_break()

    # List of Figures
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LIST OF FIGURES\n")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    figs_list = [
        ("4.1", "Overall System Architecture of EduVerse AI", "30"),
        ("4.2", "Layered Architectural Representation", "32"),
        ("4.3", "UML Use Case Diagram of EduVerse AI Platform", "35"),
        ("4.4", "Context Level Data Flow Diagram (DFD Level-0)", "36"),
        ("4.5", "Level-1 Data Flow Diagram for Student Assessment and Recommendation", "37"),
        ("4.6", "End-to-End Application Workflow Sequence", "39"),
        ("4.7", "Entity-Relationship (ER) Diagram of EduVerse AI Database", "45"),
        ("6.1", "Automated Test Execution Console Output (85 Passed)", "84"),
        ("6.2", "EduVerse AI Landing Page and Search Interface", "94"),
        ("6.3", "College Details View with Verified Official Website Button", "95"),
        ("6.4", "Side-by-Side College Comparison Matrix", "95"),
        ("6.5", "District-Wise College Distribution Analytics Dashboard", "96")
    ]
    t_lf = doc.add_table(rows=len(figs_list) + 1, cols=3)
    t_lf.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_row = t_lf.rows[0].cells
    h_row[0].paragraphs[0].add_run("Figure No.").font.bold = True
    h_row[1].paragraphs[0].add_run("Title").font.bold = True
    h_row[2].paragraphs[0].add_run("Page No.").font.bold = True
    for i, (f_no, f_title, f_pg) in enumerate(figs_list):
        row = t_lf.rows[i + 1].cells
        row[0].paragraphs[0].add_run(f_no)
        row[1].paragraphs[0].add_run(f_title)
        row[2].paragraphs[0].add_run(f_pg)
        row[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_page_break()

    # List of Abbreviations
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LIST OF ABBREVIATIONS\n")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    abbr_list = [
        ("AI", "Artificial Intelligence"),
        ("AISHE", "All India Survey on Higher Education"),
        ("API", "Application Programming Interface"),
        ("BCA", "Bachelor of Computer Applications"),
        ("B.E.", "Bachelor of Engineering"),
        ("B.Sc", "Bachelor of Science"),
        ("B.Tech", "Bachelor of Technology"),
        ("CRUD", "Create, Read, Update, Delete"),
        ("CSRF", "Cross-Site Request Forgery"),
        ("CSS", "Cascading Style Sheets"),
        ("DFD", "Data Flow Diagram"),
        ("DOM", "Document Object Model"),
        ("DOTE", "Directorate of Technical Education (Tamil Nadu)"),
        ("ER", "Entity-Relationship"),
        ("HTML", "HyperText Markup Language"),
        ("HTTP", "HyperText Transfer Protocol"),
        ("JSON", "JavaScript Object Notation"),
        ("MCA", "Master of Computer Applications"),
        ("NAAC", "National Assessment and Accreditation Council"),
        ("NBA", "National Board of Accreditation"),
        ("ORM", "Object-Relational Mapping"),
        ("REST", "Representational State Transfer"),
        ("RDBMS", "Relational Database Management System"),
        ("SaaS", "Software as a Service"),
        ("SQL", "Structured Query Language"),
        ("TNEA", "Tamil Nadu Engineering Admissions"),
        ("UI", "User Interface"),
        ("UML", "Unified Modeling Language"),
        ("URL", "Uniform Resource Locator"),
        ("UX", "User Experience"),
        ("XSS", "Cross-Site Scripting")
    ]
    t_abb = doc.add_table(rows=len(abbr_list) + 1, cols=2)
    t_abb.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_row = t_abb.rows[0].cells
    h_row[0].paragraphs[0].add_run("Abbreviation").font.bold = True
    h_row[1].paragraphs[0].add_run("Expansion").font.bold = True
    for i, (ab, exp) in enumerate(abbr_list):
        row = t_abb.rows[i + 1].cells
        row[0].paragraphs[0].add_run(ab).font.bold = True
        row[1].paragraphs[0].add_run(exp)

    doc.add_page_break()

    # ==========================================================================
    # CHAPTERS & CONTENT POPULATION
    # ==========================================================================
    md_file = BASE_DIR / "EDUVERSE_AI_PROJECT_REPORT.md"
    if md_file.exists():
        print("Reading markdown report source and formatting Word paragraphs...")
        with open(md_file, "r", encoding="utf-8") as f:
            lines = f.readlines()

        in_chapters = False
        current_table_rows = []
        in_code_block = False
        code_lines = []

        for line in lines:
            stripped = line.strip()

            if stripped.startswith("# CHAPTER 1"):
                in_chapters = True

            if not in_chapters:
                continue

            # Code blocks
            if stripped.startswith("```"):
                if in_code_block:
                    # End code block
                    in_code_block = False
                    cp = doc.add_paragraph()
                    cp.paragraph_format.line_spacing = 1.0
                    cp.paragraph_format.space_before = Pt(4)
                    cp.paragraph_format.space_after = Pt(6)
                    c_run = cp.add_run("\n".join(code_lines))
                    c_run.font.name = "Consolas"
                    c_run.font.size = Pt(9.5)
                    c_run.font.color.rgb = RGBColor(30, 41, 59)
                    code_lines = []
                else:
                    in_code_block = True
                    code_lines = []
                continue

            if in_code_block:
                code_lines.append(line.rstrip("\r\n"))
                continue

            # Heading 1
            if stripped.startswith("# CHAPTER") or stripped.startswith("# APPENDICES") or stripped.startswith("# REFERENCES"):
                doc.add_page_break()
                h1_text = stripped.lstrip("# ").strip()
                doc.add_heading(h1_text, level=1)
                continue

            # Heading 2
            if stripped.startswith("### "):
                h2_text = stripped.lstrip("# ").strip()
                doc.add_heading(h2_text, level=2)
                continue

            # Heading 3
            if stripped.startswith("#### "):
                h3_text = stripped.lstrip("# ").strip()
                doc.add_heading(h3_text, level=3)
                continue

            # Markdown Table detection
            if stripped.startswith("|") and stripped.endswith("|"):
                current_table_rows.append(stripped)
                continue
            else:
                if current_table_rows:
                    # Render table
                    rows_data = []
                    for tr in current_table_rows:
                        if "---" in tr:
                            continue
                        cells = [c.strip() for c in tr.strip("|").split("|")]
                        rows_data.append(cells)
                    current_table_rows = []

                    if rows_data:
                        cols_count = max(len(r) for r in rows_data)
                        tab = doc.add_table(rows=len(rows_data), cols=cols_count)
                        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
                        for r_idx, r_vals in enumerate(rows_data):
                            for c_idx, val in enumerate(r_vals):
                                if c_idx < cols_count:
                                    cell = tab.rows[r_idx].cells[c_idx]
                                    p = cell.paragraphs[0]
                                    p.paragraph_format.line_spacing = 1.15
                                    p.paragraph_format.space_after = Pt(2)
                                    run = p.add_run(val)
                                    run.font.name = "Times New Roman"
                                    run.font.size = Pt(10)
                                    if r_idx == 0:
                                        run.font.bold = True
                                        set_cell_background(cell, "EFF6FF")
                                    set_cell_margins(cell, 80, 80, 120, 120)
                    doc.add_paragraph() # Spacing

            # Standard paragraph or list
            if stripped:
                # Check for figure image inclusion
                if "Figure 4.1:" in stripped and (ASSETS_DIR / "fig_4_1_architecture.png").exists():
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.add_run().add_picture(str(ASSETS_DIR / "fig_4_1_architecture.png"), width=Inches(5.5))
                elif "Figure 4.3:" in stripped and (ASSETS_DIR / "fig_4_3_use_case.png").exists():
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.add_run().add_picture(str(ASSETS_DIR / "fig_4_3_use_case.png"), width=Inches(5.2))
                elif "Figure 4.4:" in stripped and (ASSETS_DIR / "fig_4_4_dfd_context.png").exists():
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.add_run().add_picture(str(ASSETS_DIR / "fig_4_4_dfd_context.png"), width=Inches(5.2))
                elif "Figure 4.7:" in stripped and (ASSETS_DIR / "fig_4_7_er_diagram.png").exists():
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.add_run().add_picture(str(ASSETS_DIR / "fig_4_7_er_diagram.png"), width=Inches(5.5))

                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                # Process basic bolding in markdown: **bold**
                parts = stripped.split("**")
                for idx, part in enumerate(parts):
                    if idx % 2 == 1:
                        r = p.add_run(part)
                        r.font.bold = True
                    else:
                        p.add_run(part)

    output_path = BASE_DIR / "EDUVERSE_AI_PROJECT_REPORT.docx"
    print(f"Saving compiled document to: {output_path}")
    doc.save(output_path)
    print("Project Report DOCX generated successfully!")


if __name__ == "__main__":
    generate_all_diagrams()
    build_academic_docx()
